"""Pull plain text out of an uploaded report so the model can read it."""

from __future__ import annotations

import csv
import io
import logging

log = logging.getLogger(__name__)

MAX_CHARS = 400_000


def extract_text(data: bytes, filename: str, content_type: str | None) -> tuple[str, str | None]:
    """
    Returns (text, error). Never raises — a file it cannot read comes
    back as empty text plus a reason the caller can show the user.
    """
    name = (filename or "").lower()

    try:
        if name.endswith(".pdf") or (content_type or "").endswith("pdf"):
            return _pdf(data), None
        if name.endswith((".xlsx", ".xlsm")):
            return _xlsx(data), None
        if name.endswith((".csv", ".tsv")):
            return _csv(data), None
        if name.endswith((".txt", ".md", ".json", ".htm", ".html")):
            return _plain(data), None
        # Unknown extension: try text, since many reports are exported as such.
        text = _plain(data)
        if text.strip():
            return text, None
        return "", (
            f"Could not read '{filename}'. Supported: PDF, XLSX, CSV, TXT, MD, HTML."
        )
    except Exception as exc:  # noqa: BLE001
        log.warning("extract failed for %s: %s", filename, exc)
        return "", f"Could not read '{filename}': {exc}"


def _plain(data: bytes) -> str:
    return data.decode("utf-8", errors="replace")[:MAX_CHARS]


def _pdf(data: bytes) -> str:
    from pypdf import PdfReader

    reader = PdfReader(io.BytesIO(data))
    parts: list[str] = []
    total = 0
    for page in reader.pages:
        try:
            t = page.extract_text() or ""
        except Exception:  # noqa: BLE001 — one bad page shouldn't lose the rest
            continue
        parts.append(t)
        total += len(t)
        if total > MAX_CHARS:
            break
    text = "\n".join(parts)
    if not text.strip():
        # Scanned PDFs carry no text layer; say so rather than returning silence.
        raise ValueError(
            "no extractable text — if this is a scanned document it needs OCR first"
        )
    return text[:MAX_CHARS]


def _xlsx(data: bytes) -> str:
    from openpyxl import load_workbook

    wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    lines: list[str] = []
    for ws in wb.worksheets:
        lines.append(f"# Sheet: {ws.title}")
        for row in ws.iter_rows(values_only=True):
            cells = [str(c) for c in row if c is not None]
            if cells:
                lines.append(" | ".join(cells))
            if len("\n".join(lines)) > MAX_CHARS:
                return "\n".join(lines)[:MAX_CHARS]
    return "\n".join(lines)[:MAX_CHARS]


def _csv(data: bytes) -> str:
    text = data.decode("utf-8", errors="replace")
    try:
        dialect = csv.Sniffer().sniff(text[:4096])
        delim = dialect.delimiter
    except csv.Error:
        delim = ","
    rows = csv.reader(io.StringIO(text), delimiter=delim)
    return "\n".join(" | ".join(r) for r in rows)[:MAX_CHARS]
